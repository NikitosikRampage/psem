"""SQLite-хранилище: просмотренные заказы, отклики, события. Экспорт в CSV/Excel."""
from __future__ import annotations

import csv
import sqlite3
import threading
from datetime import datetime, timedelta
from pathlib import Path

from .models import Order

SCHEMA = """
CREATE TABLE IF NOT EXISTS orders (
    id TEXT PRIMARY KEY,
    url TEXT, title TEXT, category TEXT,
    budget_min REAL, budget_max REAL, geo TEXT,
    client_name TEXT, client_type TEXT,
    status TEXT, reason TEXT, seen_at TEXT
);
CREATE TABLE IF NOT EXISTS responses (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    ts TEXT NOT NULL,
    order_id TEXT, title TEXT, url TEXT,
    type TEXT, price_text TEXT, price_value REAL,
    commission_percent REAL, cost REAL,
    template TEXT, message TEXT,
    status TEXT, error TEXT
);
CREATE INDEX IF NOT EXISTS responses_ts ON responses(ts);
CREATE TABLE IF NOT EXISTS events (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    ts TEXT NOT NULL, level TEXT, message TEXT
);
CREATE TABLE IF NOT EXISTS kv (key TEXT PRIMARY KEY, value TEXT);
"""

RESPONSE_COLUMNS = [
    "id", "ts", "order_id", "title", "url", "type", "price_text", "price_value",
    "commission_percent", "cost", "template", "message", "status", "error",
]


def _ts(dt: datetime | None = None) -> str:
    return (dt or datetime.now()).isoformat(timespec="seconds")


class Storage:
    def __init__(self, path: Path | str):
        path = Path(path)
        if str(path) != ":memory:":
            path.parent.mkdir(parents=True, exist_ok=True)
        self._lock = threading.RLock()
        self._conn = sqlite3.connect(str(path), check_same_thread=False)
        self._conn.row_factory = sqlite3.Row
        with self._lock:
            self._conn.executescript(SCHEMA)
            self._conn.commit()

    def close(self) -> None:
        with self._lock:
            self._conn.close()

    def _exec(self, sql: str, params=()) -> sqlite3.Cursor:
        with self._lock:
            cur = self._conn.execute(sql, params)
            self._conn.commit()
            return cur

    def _query(self, sql: str, params=()) -> list[sqlite3.Row]:
        with self._lock:
            return self._conn.execute(sql, params).fetchall()

    # --- заказы ---
    def is_seen(self, order_id: str) -> bool:
        # «Устаревшие» по ленте заказы проверяются снова: клиент может обновить заказ.
        return bool(self._query("SELECT 1 FROM orders WHERE id=? AND status!='too_old'", (order_id,)))

    def mark_order(self, order: Order, status: str, reason: str = "") -> None:
        self._exec(
            """INSERT INTO orders (id,url,title,category,budget_min,budget_max,geo,client_name,
                                   client_type,status,reason,seen_at)
               VALUES (?,?,?,?,?,?,?,?,?,?,?,?)
               ON CONFLICT(id) DO UPDATE SET status=excluded.status, reason=excluded.reason""",
            (order.id, order.url, order.title, order.category, order.budget_min, order.budget_max,
             order.geo, order.client_name, order.client_type, status, reason, _ts()),
        )

    def forget_dry_run_orders(self) -> int:
        """Заказы, обработанные в тестовом режиме, снова станут доступны для реальных откликов."""
        return self._exec("DELETE FROM orders WHERE status='dry_run'").rowcount

    # --- отклики ---
    def add_response(self, **fields) -> int:
        fields.setdefault("ts", _ts())
        cols = [c for c in RESPONSE_COLUMNS if c != "id" and c in fields]
        cur = self._exec(
            f"INSERT INTO responses ({','.join(cols)}) VALUES ({','.join('?' * len(cols))})",
            [fields[c] for c in cols],
        )
        return cur.lastrowid

    def recent_responses(self, limit: int = 200) -> list[dict]:
        return [dict(r) for r in self._query("SELECT * FROM responses ORDER BY id DESC LIMIT ?", (limit,))]

    def count_sent(self, since: datetime, type_: str | None = None) -> int:
        sql = "SELECT COUNT(*) FROM responses WHERE status='sent' AND ts>=?"
        params: list = [_ts(since)]
        if type_:
            sql += " AND type=?"
            params.append(type_)
        return self._query(sql, params)[0][0]

    def sum_cost(self, since: datetime) -> float:
        row = self._query(
            "SELECT COALESCE(SUM(cost),0) FROM responses WHERE status='sent' AND type='paid' AND ts>=?",
            (_ts(since),),
        )
        return float(row[0][0])

    def sent_timestamps(self, since: datetime) -> list[datetime]:
        rows = self._query("SELECT ts FROM responses WHERE status='sent' AND ts>=? ORDER BY ts", (_ts(since),))
        return [datetime.fromisoformat(r[0]) for r in rows]

    # --- события и kv ---
    def add_event(self, level: str, message: str) -> None:
        self._exec("INSERT INTO events (ts,level,message) VALUES (?,?,?)", (_ts(), level, message))

    def recent_events(self, limit: int = 300) -> list[dict]:
        rows = self._query("SELECT * FROM events ORDER BY id DESC LIMIT ?", (limit,))
        return [dict(r) for r in reversed(rows)]

    def get_kv(self, key: str, default: str | None = None) -> str | None:
        rows = self._query("SELECT value FROM kv WHERE key=?", (key,))
        return rows[0][0] if rows else default

    def set_kv(self, key: str, value) -> None:
        self._exec(
            "INSERT INTO kv (key,value) VALUES (?,?) ON CONFLICT(key) DO UPDATE SET value=excluded.value",
            (key, str(value)),
        )

    def next_counter(self, key: str) -> int:
        with self._lock:
            value = int(self.get_kv(key, "0"))
            self.set_kv(key, value + 1)
            return value

    # --- статистика и экспорт ---
    def stats(self, now: datetime | None = None) -> dict:
        now = now or datetime.now()
        day_start = now.replace(hour=0, minute=0, second=0, microsecond=0)

        hour_ago = now - timedelta(hours=1)
        by_status = {
            r[0]: r[1]
            for r in self._query(
                "SELECT status, COUNT(*) FROM responses WHERE ts>=? GROUP BY status", (_ts(day_start),)
            )
        }
        return {
            "sent_today": self.count_sent(day_start),
            "sent_hour": self.count_sent(hour_ago),
            "paid_today": self.count_sent(day_start, "paid"),
            "commission_today": self.count_sent(day_start, "commission"),
            "spent_today": self.sum_cost(day_start),
            "dry_run_today": by_status.get("dry_run", 0),
            "errors_today": by_status.get("error", 0),
            "orders_seen": self._query("SELECT COUNT(*) FROM orders")[0][0],
            "orders_skipped": self._query("SELECT COUNT(*) FROM orders WHERE status='skipped'")[0][0],
            "total_sent": self._query("SELECT COUNT(*) FROM responses WHERE status='sent'")[0][0],
        }

    def daily_summary(self) -> list[dict]:
        rows = self._query(
            """SELECT substr(ts,1,10) AS day,
                      SUM(status='sent') AS sent,
                      SUM(status='sent' AND type='paid') AS paid,
                      SUM(status='sent' AND type='commission') AS commission,
                      COALESCE(SUM(CASE WHEN status='sent' AND type='paid' THEN cost END),0) AS spent,
                      SUM(status='dry_run') AS dry_run,
                      SUM(status='error') AS errors
               FROM responses GROUP BY day ORDER BY day"""
        )
        return [dict(r) for r in rows]

    def export_csv(self, path: Path | str) -> Path:
        path = Path(path)
        path.parent.mkdir(parents=True, exist_ok=True)
        rows = self._query("SELECT * FROM responses ORDER BY id")
        with path.open("w", newline="", encoding="utf-8-sig") as fh:
            writer = csv.writer(fh, delimiter=";")
            writer.writerow(RESPONSE_COLUMNS)
            for r in rows:
                writer.writerow([r[c] for c in RESPONSE_COLUMNS])
        return path

    def export_xlsx(self, path: Path | str) -> Path:
        from openpyxl import Workbook

        path = Path(path)
        path.parent.mkdir(parents=True, exist_ok=True)
        wb = Workbook()
        ws = wb.active
        ws.title = "Отклики"
        ws.append(RESPONSE_COLUMNS)
        for r in self._query("SELECT * FROM responses ORDER BY id"):
            ws.append([r[c] for c in RESPONSE_COLUMNS])

        ws2 = wb.create_sheet("По дням")
        summary = self.daily_summary()
        ws2.append(["day", "sent", "paid", "commission", "spent", "dry_run", "errors"])
        for row in summary:
            ws2.append([row[k] for k in ("day", "sent", "paid", "commission", "spent", "dry_run", "errors")])

        ws3 = wb.create_sheet("Заказы")
        cols = ["id", "seen_at", "status", "reason", "title", "category", "budget_min", "budget_max",
                "geo", "client_name", "client_type", "url"]
        ws3.append(cols)
        for r in self._query("SELECT * FROM orders ORDER BY seen_at"):
            ws3.append([r[c] for c in cols])
        wb.save(path)
        return path

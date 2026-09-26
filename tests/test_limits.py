from datetime import datetime, timedelta

from profi_bot.limits import Limits
from profi_bot.storage import Storage

NOW = datetime(2026, 9, 26, 14, 30)


def make(sent):
    s = Storage(":memory:")
    for ts, type_, cost in sent:
        s.add_response(ts=ts.isoformat(timespec="seconds"), type=type_, cost=cost, status="sent")
    s.add_response(ts=NOW.isoformat(), type="paid", cost=999, status="dry_run")  # не считается
    return s, Limits(s)


def test_per_day_and_hour():
    s, lim = make([(NOW - timedelta(hours=5), "paid", 50)] * 2 + [(NOW - timedelta(minutes=m), "commission", 0)
                                                                   for m in (10, 20, 50)])
    assert not lim.check_global({"per_day": 10, "per_hour": 4}, NOW).blocked
    st = lim.check_global({"per_day": 5, "per_hour": 0}, NOW)
    assert st.blocked and st.resume_at == datetime(2026, 9, 27)
    st = lim.check_global({"per_day": 0, "per_hour": 3}, NOW)
    assert st.blocked and st.resume_at == NOW - timedelta(minutes=50) + timedelta(hours=1)
    st = lim.check_global({"per_day": 0, "per_hour": 2}, NOW)
    assert st.resume_at == NOW - timedelta(minutes=20) + timedelta(hours=1)


def test_paid_limits():
    s, lim = make([(NOW - timedelta(hours=1), "paid", 300), (NOW - timedelta(hours=2), "paid", 300)])
    assert lim.paid_allowed({"paid_per_day": 3, "budget_per_day": 1000}, 300, NOW)
    assert not lim.paid_allowed({"paid_per_day": 2, "budget_per_day": 0}, 10, NOW)
    assert not lim.paid_allowed({"paid_per_day": 0, "budget_per_day": 800}, 300, NOW)
    assert not lim.paid_allowed({"budget_per_day": 600}, 0, NOW)
    cfg = {"paid_per_day": 2}
    assert not lim.check_global(cfg, NOW, commission_allowed=True).blocked
    assert lim.check_global(cfg, NOW, commission_allowed=False).blocked


def test_yesterday_not_counted():
    s, lim = make([(NOW - timedelta(days=1), "paid", 500)] * 5)
    assert not lim.check_global({"per_day": 1}, NOW).blocked
    assert lim.paid_allowed({"paid_per_day": 1, "budget_per_day": 100}, 100, NOW)


def test_stats_and_export(tmp_path):
    s, _ = make([(datetime.now(), "paid", 100), (datetime.now(), "commission", 0)])
    st = s.stats()
    assert st["sent_today"] == 2 and st["paid_today"] == 1 and st["spent_today"] == 100
    assert st["dry_run_today"] in (0, 1)
    csv_path = s.export_csv(tmp_path / "r.csv")
    assert csv_path.read_text(encoding="utf-8-sig").count("\n") == 4
    from openpyxl import load_workbook

    wb = load_workbook(s.export_xlsx(tmp_path / "r.xlsx"))
    assert wb.sheetnames == ["Отклики", "По дням", "Заказы"] and wb["Отклики"].max_row == 4
